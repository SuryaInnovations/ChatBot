"""
tools/file_parser.py — Multi-format Document and Spreadsheet Parser.

Extracts structured text and metadata from:
- PDF (page numbers, text per page)
- DOCX (paragraphs, tables, headings)
- PPTX (slides, shape text)
- XLSX / CSV (sheet names, tabular rows, column headers)
- TXT / MD / JSON / Code files
"""

import os
import csv
import json
import logging
from typing import List, Dict, Any, Optional

logger = logging.getLogger("chatbot.tools.file_parser")

MAX_FILE_BYTES = 50 * 1024 * 1024  # 50 MB limit


class ParsedSection:
    def __init__(
        self,
        text: str,
        file_name: str,
        page: Optional[int] = None,
        sheet_name: Optional[str] = None,
        section_type: str = "text"
    ):
        self.text = text.strip()
        self.file_name = file_name
        self.page = page
        self.sheet_name = sheet_name
        self.section_type = section_type
        self.char_count = len(self.text)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "text": self.text,
            "file_name": self.file_name,
            "page": self.page,
            "sheet_name": self.sheet_name,
            "section_type": self.section_type,
            "char_count": self.char_count
        }


def parse_document(file_path: str, file_name: Optional[str] = None) -> List[Dict[str, Any]]:
    """
    Parses a document file and returns a list of section dictionaries with location metadata.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    size = os.path.getsize(file_path)
    if size > MAX_FILE_BYTES:
        raise ValueError(f"File size ({round(size / 1024 / 1024, 2)} MB) exceeds maximum allowed 50 MB.")

    fname = file_name or os.path.basename(file_path)
    ext = os.path.splitext(file_path)[1].lower()
    sections: List[ParsedSection] = []

    try:
        if ext == ".pdf":
            # Try pypdf or fitz (PyMuPDF)
            try:
                import fitz # PyMuPDF
                doc = fitz.open(file_path)
                for i, page in enumerate(doc, start=1):
                    txt = page.get_text()
                    if txt.strip():
                        sections.append(ParsedSection(txt, fname, page=i, section_type="page"))
            except ImportError:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                for i, page in enumerate(reader.pages, start=1):
                    txt = page.extract_text() or ""
                    if txt.strip():
                        sections.append(ParsedSection(txt, fname, page=i, section_type="page"))

        elif ext == ".docx":
            import docx
            doc = docx.Document(file_path)
            full_text = []
            for p in doc.paragraphs:
                if p.text.strip():
                    full_text.append(p.text.strip())
            for t in doc.tables:
                for row in t.rows:
                    row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_txt:
                        full_text.append(row_txt)
            if full_text:
                sections.append(ParsedSection("\n\n".join(full_text), fname, page=1, section_type="document"))

        elif ext in (".pptx", ".ppt"):
            from pptx import Presentation
            prs = Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides, start=1):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                if slide_texts:
                    sections.append(ParsedSection("\n".join(slide_texts), fname, page=slide_idx, section_type="slide"))

        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                rows_text = []
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if row_vals:
                        rows_text.append(" | ".join(row_vals))
                if rows_text:
                    sections.append(ParsedSection("\n".join(rows_text[:500]), fname, sheet_name=sheetname, section_type="sheet"))

        elif ext == ".csv":
            import pandas as pd
            df = pd.read_csv(file_path, nrows=500)
            sections.append(ParsedSection(df.to_string(index=False), fname, sheet_name="csv_data", section_type="table"))

        elif ext in (".json", ".jsonl"):
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read(500000)
            sections.append(ParsedSection(content, fname, page=1, section_type="json"))

        else:
            # Plain text, markdown, source code
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read(500000)
            sections.append(ParsedSection(content, fname, page=1, section_type="text"))

    except Exception as e:
        logger.error(f"Error parsing file {fname}: {e}")
        raise RuntimeError(f"Failed to parse document '{fname}': {str(e)}")

    return [s.to_dict() for s in sections]

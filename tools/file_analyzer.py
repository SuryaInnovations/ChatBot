"""
File intelligence and analysis tool supporting CSV, Excel, PDF, DOCX, TXT, JSON, and code files.
"""

import os
from typing import Dict, Any

def analyze_file(file_path: str, max_chars: int = 5000) -> Dict[str, Any]:
    """
    Parses and extracts key structural information, data summary, or text from a file.
    """
    if not os.path.exists(file_path):
        return {"success": False, "error": f"File not found: {file_path}"}

    ext = os.path.splitext(file_path)[1].lower()
    filename = os.path.basename(file_path)
    file_size_kb = round(os.path.getsize(file_path) / 1024, 2)

    try:
        if ext in [".csv", ".tsv"]:
            import pandas as pd
            sep = "\t" if ext == ".tsv" else ","
            df = pd.read_csv(file_path, sep=sep)
            info = {
                "success": True,
                "file_type": "CSV/Tabular",
                "filename": filename,
                "file_size_kb": file_size_kb,
                "rows": len(df),
                "columns": list(df.columns),
                "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()},
                "missing_values": df.isnull().sum().to_dict(),
                "head": df.head(5).to_dict(orient="records"),
                "summary_stats": df.describe(include="all").to_dict() if len(df) > 0 else {}
            }
            return info

        elif ext in [".xlsx", ".xls"]:
            import pandas as pd
            excel_file = pd.ExcelFile(file_path)
            sheets_data = {}
            for sheet in excel_file.sheet_names[:3]: # inspect first 3 sheets
                df = pd.read_excel(file_path, sheet_name=sheet)
                sheets_data[sheet] = {
                    "rows": len(df),
                    "columns": list(df.columns),
                    "head": df.head(3).to_dict(orient="records")
                }
            return {
                "success": True,
                "file_type": "Excel",
                "filename": filename,
                "file_size_kb": file_size_kb,
                "sheet_names": excel_file.sheet_names,
                "sheets": sheets_data
            }

        elif ext == ".pdf":
            import pypdf
            reader = pypdf.PdfReader(file_path)
            num_pages = len(reader.pages)
            extracted_text = []
            for i in range(min(num_pages, 10)): # extract up to first 10 pages
                page_text = reader.pages[i].extract_text() or ""
                extracted_text.append(f"--- Page {i+1} ---\n{page_text.strip()}")
            full_text = "\n\n".join(extracted_text)
            return {
                "success": True,
                "file_type": "PDF",
                "filename": filename,
                "file_size_kb": file_size_kb,
                "total_pages": num_pages,
                "preview_text": full_text[:max_chars],
                "truncated": len(full_text) > max_chars
            }

        elif ext in [".docx", ".doc"]:
            import docx
            doc = docx.Document(file_path)
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            tables_count = len(doc.tables)
            full_text = "\n\n".join(paras)
            return {
                "success": True,
                "file_type": "Word Document",
                "filename": filename,
                "file_size_kb": file_size_kb,
                "paragraph_count": len(paras),
                "tables_count": tables_count,
                "preview_text": full_text[:max_chars],
                "truncated": len(full_text) > max_chars
            }

        elif ext in [".txt", ".md", ".json", ".py", ".js", ".html", ".css", ".sql", ".yaml", ".yml", ".xml", ".sh", ".ps1"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            line_count = len(content.splitlines())
            return {
                "success": True,
                "file_type": "Text/Code",
                "filename": filename,
                "file_size_kb": file_size_kb,
                "lines": line_count,
                "content": content[:max_chars],
                "truncated": len(content) > max_chars
            }
        else:
            return {
                "success": True,
                "file_type": "Binary/Other",
                "filename": filename,
                "file_size_kb": file_size_kb,
                "message": f"File format {ext} metadata retrieved, content binary or unparsed."
            }

    except Exception as e:
        return {"success": False, "error": f"Failed analyzing {filename}: {str(e)}"}
